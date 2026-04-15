#!/usr/bin/env python3
"""
GDPVal judge — scores deliverables against a rubric using an LLM.

Provider chain (first available wins):
  1. Gemini        (GEMINI_API_KEY)           — paper default, multimodal PDF.
  2. Anthropic     (ANTHROPIC_API_KEY / CLAUDE_CODE_OAUTH_TOKEN) — multimodal PDF.
  3. Ollama Cloud  (http://host.docker.internal:11434) — free, text-only.
  4. Fallback      (no judge reachable) → reward 0.5.

Override with JUDGE_PROVIDER=gemini|anthropic|ollama and JUDGE_MODEL=<id>.
"""

from __future__ import annotations

import base64
import json
import os
import subprocess
import sys
import urllib.request
from abc import ABC, abstractmethod
from pathlib import Path

RUBRIC_PATH = Path("/tests/rubric.json")
DELIVERABLES_PATH = Path("/tests/deliverables.txt")
REWARD_PATH = Path("/logs/verifier/reward.json")
PDF_CACHE_DIR = Path("/tmp/judge-pdfs")

MAX_CONTENT_CHARS = 60_000
MAX_INLINE_BYTES = 18 * 1024 * 1024  # Gemini/Anthropic inline cap ~20MB
PDF_CONVERTIBLE = {".xlsx", ".xlsm", ".xls", ".docx", ".doc",
                   ".pptx", ".ppt", ".odt", ".ods"}

INSTRUCTIONS = (
    "You are an expert professional grader for the GDPVal benchmark.\n"
    "Score the deliverables provided below against the rubric.\n\n"
    "Score on a scale from 0.0 to 1.0 based on the fraction of rubric points earned.\n"
    'Respond with ONLY a JSON object: {"score": <float 0-1>, "reason": "<brief>"}\n'
)


# ---------- Text extraction ----------

def _extract_xlsx(path: Path) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True, read_only=True)
    parts: list[str] = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        parts.append(f"### Sheet: {sheet_name}")
        for row in ws.iter_rows(values_only=True):
            parts.append("\t".join("" if c is None else str(c) for c in row))
    return "\n".join(parts)


def _extract_pdf(path: Path) -> str:
    import pdfplumber
    with pdfplumber.open(path) as pdf:
        return "\n\n".join((page.extract_text() or "") for page in pdf.pages)


def _extract_docx(path: Path) -> str:
    from docx import Document
    doc = Document(path)
    return "\n".join(p.text for p in doc.paragraphs)


_TEXT_EXTRACTORS = {
    ".xlsx": _extract_xlsx, ".xlsm": _extract_xlsx, ".xls": _extract_xlsx,
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
}


def extract_text(path: Path) -> str:
    ext = path.suffix.lower()
    try:
        content = _TEXT_EXTRACTORS[ext](path) if ext in _TEXT_EXTRACTORS \
            else path.read_text(encoding="utf-8", errors="replace")
    except Exception as e:
        return f"[extraction failed for {path.name}: {e}]"
    if len(content) > MAX_CONTENT_CHARS:
        content = content[:MAX_CONTENT_CHARS] + "\n[...truncated...]"
    return content


# ---------- Attachment ----------

class Attachment:
    """A deliverable file exposed either as inline PDF bytes or extracted text."""

    def __init__(self, name: str, path: Path):
        self.name = name
        self.path = path

    def as_pdf_or_text(self) -> tuple[str, str]:
        """Returns ('pdf', base64) when PDF conversion succeeds and fits inline, else ('text', str)."""
        pdf = self._convert_to_pdf()
        if pdf is not None:
            b64 = self._inline_b64(pdf)
            if b64 is not None:
                return "pdf", b64
        return "text", extract_text(self.path)

    def as_text(self) -> str:
        return extract_text(self.path)

    def _convert_to_pdf(self) -> Path | None:
        ext = self.path.suffix.lower()
        if ext == ".pdf":
            return self.path
        if ext not in PDF_CONVERTIBLE:
            return None
        PDF_CACHE_DIR.mkdir(parents=True, exist_ok=True)
        target = PDF_CACHE_DIR / (self.path.stem + ".pdf")
        if target.exists():
            return target
        try:
            subprocess.run(
                ["libreoffice", "--headless", "--convert-to", "pdf",
                 "--outdir", str(PDF_CACHE_DIR), str(self.path)],
                check=True, capture_output=True, timeout=120,
            )
        except Exception as e:
            print(f"PDF conversion failed for {self.path.name}: {e}", file=sys.stderr)
            return None
        return target if target.exists() else None

    @staticmethod
    def _inline_b64(path: Path) -> str | None:
        if path.stat().st_size > MAX_INLINE_BYTES:
            return None
        return base64.b64encode(path.read_bytes()).decode()


# ---------- HTTP ----------

def http_post_json(url: str, body: dict, headers: dict, timeout: int = 300) -> dict:
    data = json.dumps(body).encode()
    req = urllib.request.Request(
        url, data=data,
        headers={"Content-Type": "application/json", **headers},
    )
    with urllib.request.urlopen(req, timeout=timeout) as resp:
        return json.loads(resp.read())


def parse_score(text: str) -> dict:
    start = text.find("{")
    end = text.rfind("}") + 1
    return json.loads(text[start:end])


# ---------- Judges ----------

class Judge(ABC):
    """Abstract judge: checks availability, sends rubric+deliverables, returns {score, reason}."""

    name: str = ""
    default_model: str = ""

    def __init__(self):
        self.model = os.environ.get("JUDGE_MODEL", self.default_model)

    @abstractmethod
    def available(self) -> bool: ...

    @abstractmethod
    def score(self, rubric: str, attachments: list[Attachment]) -> dict: ...

    @staticmethod
    def _prompt(rubric: str) -> str:
        return f"{INSTRUCTIONS}\nRUBRIC:\n{rubric}\n"


class GeminiJudge(Judge):
    name = "gemini"
    default_model = "gemini-3.1-pro-preview"

    def available(self) -> bool:
        return bool(os.environ.get("GEMINI_API_KEY"))

    def score(self, rubric: str, attachments: list[Attachment]) -> dict:
        url = (
            "https://generativelanguage.googleapis.com/v1beta/models/"
            f"{self.model}:generateContent?key={os.environ['GEMINI_API_KEY']}"
        )
        parts: list[dict] = [{"text": self._prompt(rubric)}]
        for att in attachments:
            kind, payload = att.as_pdf_or_text()
            if kind == "pdf":
                parts.append({"text": f"\n--- {att.name} (attached PDF) ---"})
                parts.append({"inline_data": {"mime_type": "application/pdf", "data": payload}})
            else:
                parts.append({"text": f"\n--- {att.name} ---\n{payload}"})

        data = http_post_json(url, {
            "contents": [{"parts": parts}],
            "generationConfig": {"temperature": 0.0, "maxOutputTokens": 500},
        }, {})
        return parse_score(data["candidates"][0]["content"]["parts"][0]["text"])


class AnthropicJudge(Judge):
    name = "anthropic"
    default_model = "claude-sonnet-4-5"

    def _headers(self) -> dict | None:
        if key := os.environ.get("ANTHROPIC_API_KEY"):
            return {"x-api-key": key, "anthropic-version": "2023-06-01"}
        if oauth := os.environ.get("CLAUDE_CODE_OAUTH_TOKEN"):
            return {
                "Authorization": f"Bearer {oauth}",
                "anthropic-version": "2023-06-01",
                "anthropic-beta": "oauth-2025-04-20",
            }
        return None

    def available(self) -> bool:
        return self._headers() is not None

    def score(self, rubric: str, attachments: list[Attachment]) -> dict:
        content: list[dict] = [{"type": "text", "text": self._prompt(rubric)}]
        for att in attachments:
            kind, payload = att.as_pdf_or_text()
            if kind == "pdf":
                content.append({"type": "text", "text": f"\n--- {att.name} (attached PDF) ---"})
                content.append({
                    "type": "document",
                    "source": {"type": "base64", "media_type": "application/pdf", "data": payload},
                })
            else:
                content.append({"type": "text", "text": f"\n--- {att.name} ---\n{payload}"})

        data = http_post_json("https://api.anthropic.com/v1/messages", {
            "model": self.model, "max_tokens": 500, "temperature": 0.0,
            "messages": [{"role": "user", "content": content}],
        }, self._headers())
        return parse_score(data["content"][0]["text"])


class OllamaJudge(Judge):
    name = "ollama"
    default_model = "gpt-oss:120b-cloud"

    def available(self) -> bool:
        return True  # free fallback; if unreachable, score() raises and chain moves on

    def score(self, rubric: str, attachments: list[Attachment]) -> dict:
        host = os.environ.get("OLLAMA_HOST", "http://host.docker.internal:11434")
        body = self._prompt(rubric) + "\nDELIVERABLE CONTENTS:\n" + "".join(
            f"\n--- {a.name} ---\n{a.as_text()}" for a in attachments
        )
        data = http_post_json(f"{host}/api/generate", {
            "model": self.model, "prompt": body, "stream": False,
            "options": {"temperature": 0.0},
        }, {}, timeout=300)
        return parse_score(data.get("response", ""))


JUDGE_REGISTRY: dict[str, type[Judge]] = {
    GeminiJudge.name: GeminiJudge,
    AnthropicJudge.name: AnthropicJudge,
    OllamaJudge.name: OllamaJudge,
}


def judge_chain() -> list[Judge]:
    forced = os.environ.get("JUDGE_PROVIDER")
    if forced and forced in JUDGE_REGISTRY:
        return [JUDGE_REGISTRY[forced]()]
    return [cls() for cls in JUDGE_REGISTRY.values()]


# ---------- Deliverable discovery & output ----------

def find_deliverable(name: str) -> Path | None:
    for root in (Path("/app/output"), Path("/app")):
        candidate = root / name
        if candidate.is_file():
            return candidate
    return None


def write_reward(reward: float, reason: str, judge: str = "none") -> None:
    REWARD_PATH.parent.mkdir(parents=True, exist_ok=True)
    REWARD_PATH.write_text(json.dumps({"reward": reward}))
    (REWARD_PATH.parent / "judge.json").write_text(json.dumps({
        "reward": reward, "reason": reason, "judge": judge,
    }, indent=2))


def main() -> None:
    deliverables = [
        d.strip() for d in DELIVERABLES_PATH.read_text().splitlines() if d.strip()
    ]

    attachments: list[Attachment] = []
    missing: list[str] = []
    for name in deliverables:
        path = find_deliverable(name)
        if path is None:
            missing.append(name)
        else:
            attachments.append(Attachment(name, path))

    if missing:
        write_reward(0.0, f"missing deliverables: {', '.join(missing)}")
        print(f"Missing: {missing}. Reward: 0.0")
        return

    rubric = RUBRIC_PATH.read_text() if RUBRIC_PATH.exists() else "[]"

    for judge in judge_chain():
        if not judge.available():
            continue
        try:
            result = judge.score(rubric, attachments)
        except Exception as e:
            print(f"Judge {judge.name} errored: {e}", file=sys.stderr)
            continue
        score = max(0.0, min(1.0, float(result.get("score", 0.5))))
        reason = str(result.get("reason", ""))[:500]
        write_reward(score, reason, judge=judge.name)
        print(f"Judge: {judge.name}. Reward: {score}. Reason: {reason}")
        return

    write_reward(0.5, "deliverables exist but no judge reachable", judge="fallback")
    print("No judge reachable. Reward: 0.5")


if __name__ == "__main__":
    main()
